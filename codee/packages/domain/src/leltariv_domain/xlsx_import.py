from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

EXCEL_DATE_BASE = date(1899, 12, 30)


def clean_header(value: object) -> str:
    """Az SAP-export fejlécének elejéről-végéről leveszi a szóközöket."""
    if value is None:
        return ""

    return str(value).strip()


def as_source_text(value: object) -> str | None:
    """Forrásértéket szöveggé alakít anélkül, hogy a meglévő 0-k eltűnnének."""
    if value is None:
        return None

    if isinstance(value, str):
        cleaned = value.strip()
        return cleaned or None

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip() or None


def excel_serial_to_date(value: object) -> date | None:
    """Excel-sorszámot Python dátummá alakít a 1899-12-30 bázissal."""
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, bool):
        raise ValueError("Az Excel-dátum nem lehet logikai érték.")

    if isinstance(value, (int, float)):
        return EXCEL_DATE_BASE + timedelta(days=int(value))

    raise ValueError(f"Nem értelmezhető Excel-dátum: {value!r}")


def zone_code_from_site(value: object) -> str:
    """A Telephely első három karakteréből képezi a körzetkódot."""
    site_code = as_source_text(value)

    if site_code is None or len(site_code) < 3:
        raise ValueError(f"Érvénytelen telephely: {value!r}")

    return site_code[:3]


def is_asset_row(row: tuple[object, ...]) -> bool:
    """Az SAP-összesítő soroknál üres az Alszám, ezek nem eszköztételek."""
    return len(row) > 1 and row[1] not in (None, "")


def read_asset_rows(path: Path) -> Iterator[dict[str, object]]:
    """XLSX-ből csak valódi eszközsorokat ad vissza, fejlécnév szerinti dictként."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active

    if worksheet is None:
        workbook.close()
        raise ValueError(f"Nincs aktív munkalap: {path}")

    try:
        row_iterator = worksheet.iter_rows(values_only=True)
        raw_headers = next(row_iterator, None)

        if raw_headers is None:
            raise ValueError(f"Üres munkalap: {path}")

        headers = [clean_header(header) for header in raw_headers]

        for row in row_iterator:
            if not is_asset_row(row):
                continue

            yield {
                header: value
                for header, value in zip(headers, row, strict=False)
                if header
            }
    finally:
        workbook.close()